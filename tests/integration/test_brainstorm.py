"""Independent application sessions use native Teams and same-owner context."""

from __future__ import annotations

from typing import Any

import pytest
from langchain_core.messages import HumanMessage

from octop.infra.agents.teams.brainstorm import service_for
from octop.infra.errors import OctopError
from tests.integration.test_applications import grant_product
from tests.integration.test_applications import products as products


# Reuse the native platform fixture; do not invoke external LLMs.
@pytest.fixture
async def advisory(products: Any) -> Any:
    client, server, admin, *_ = products
    result = await client.post(
        "/api/admin/applications",
        headers=admin,
        json={
            "application_id": "super-advisory",
            "name": "半人马AI-超级参谋团",
            "capabilities": ["chat", "teams", "brainstorm", "zhijun_context"],
            "agent_template": {"default_model": "openai/gpt-4o"},
            "resource_requirements": {"default_model": {"kind": "model"}},
        },
    )
    assert result.status_code == 201, result.text
    auth = await grant_product(client, server, admin, "reader", "super-advisory")
    yield client, server, admin, auth


@pytest.fixture
async def advisory_native(tmp_octop_home: Any) -> Any:
    from tests.support.app import octop_client
    from tests.support.auth import auth_header, bootstrap_admin, create_user, seed_openai_provider

    async with octop_client(tmp_octop_home, patch_llm=False) as (client, server):
        await bootstrap_admin(client, tmp_octop_home)
        admin = await auth_header(client)
        await seed_openai_provider(client, admin)
        await create_user(client, admin, username="reader")
        created = await client.post(
            "/api/admin/applications",
            headers=admin,
            json={
                "application_id": "super-advisory",
                "name": "超级参谋团",
                "capabilities": ["chat", "teams", "brainstorm", "zhijun_context"],
                "agent_template": {"default_model": "openai/gpt-4o"},
            },
        )
        assert created.status_code == 201
        auth = await grant_product(client, server, admin, "reader", "super-advisory")
        yield client, server, admin, auth


async def confirm_agenda(client: Any, registry: Any, auth: Any, state: Any) -> Any:
    from octop.infra.agents.teams.brainstorm import build_brainstorm_tools

    checkpoint = next(
        t
        for t in build_brainstorm_tools(registry, state["team_id"])
        if t.name == "brainstorm_checkpoint"
    )
    await checkpoint.ainvoke(
        {
            "summary": "先按用户目标比较方案",
            "ideas": [],
            "directions": [],
            "questions": [],
            "phase": "agenda",
            "goal": state["title"],
            "constraints": ["时间有限"],
            "agenda": "比较两条最小验证路线",
        },
        config={
            "configurable": {
                "thread_id": state["thread_id"],
                "agent_id": state["team_id"],
                "user": str(state["user_id"]),
            }
        },
    )
    current = (
        await client.get(f"/api/brainstorm/sessions/{state['thread_id']}", headers=auth)
    ).json()
    response = await client.post(
        f"/api/brainstorm/sessions/{state['thread_id']}/control",
        headers=auth,
        json={
            "action": "start",
            "decision_id": current["decision"]["id"],
            "revision": current["revision"],
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


async def test_profile_is_idempotent_five_private_native_agents(advisory: Any) -> None:
    client, server, admin, auth = advisory
    first = await client.post("/api/brainstorm/profile", headers=auth, json={})
    assert first.status_code == 200, first.text
    profile = first.json()
    assert len(profile["members"]) == 4
    assert len({member["agent_id"] for member in profile["members"]}) == 4
    generic = await client.post("/api/applications/super-advisory/profile", headers=auth)
    assert generic.status_code == 201, generic.text
    assert generic.json()["agent_id"] == profile["team_id"]
    schema = client._octop_app.openapi()
    operations = [
        op
        for path, methods in schema["paths"].items()
        if path.startswith("/api/brainstorm")
        for op in methods.values()
    ]
    assert len(operations) == 9
    assert all(
        op["summary"]
        and next(r for code, r in op["responses"].items() if code.startswith("2"))["content"]
        for op in operations
    )
    from copy import copy
    from dataclasses import replace

    import httpx

    from octop.api.app import build_app

    docs_server = copy(server)
    docs_server.services = replace(
        server.services, config=replace(server.services.config, enable_api_docs=True)
    )
    docs_app = build_app(docs_server)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=docs_app), base_url="http://testserver"
    ) as docs_client:
        assert (await docs_client.get("/api/docs", headers=admin)).status_code == 200
    second = await client.post("/api/brainstorm/profile", headers=auth, json={})
    assert second.json()["team_id"] == profile["team_id"]
    assert second.json()["members"] == profile["members"]
    registry = server.app_runtime.agent_registry
    host = registry.get_row(profile["team_id"])
    assert host.kind == "team" and host.application_id == "super-advisory"
    assert registry._build_harness_config(host).team_enabled
    assert set(registry.teams.member_ids(host.agent_id)) == {
        m["agent_id"] for m in profile["members"]
    }
    assert (await client.get(f"/api/agents/{host.agent_id}", headers=admin)).status_code == 404
    other = await grant_product(client, server, admin, "builder", "super-advisory")
    assert (await client.get(f"/api/agents/{host.agent_id}", headers=other)).status_code == 403
    assert (await client.get("/api/zhijun/profile", headers=auth)).status_code == 404


async def test_pause_restart_and_owned_history_sources(advisory: Any, monkeypatch: Any) -> None:
    client, server, admin, auth = advisory
    registry = server.app_runtime.agent_registry
    service = service_for(registry)
    monkeypatch.setattr(service, "_launch", lambda *args: None)
    await client.post("/api/brainstorm/profile", headers=auth, json={})
    user = server.user_manager.get("reader")
    source = server.services.agent_repo.create(
        agent_id="my-zhijun", user_id=user.id, name="知君", application_id="zhijun"
    )
    source_thread = server.app_runtime.gateway.thread_registry.create_thread(
        agent_id=source,
        user_id=user.id,
        channel_type="dashboard",
        session_key="my-zhijun:dashboard:test",
        title="咖啡店计划",
    )
    from octop.infra.history.projection import message_inputs

    repo = server.services.thread_message_repo
    repo.mark_projection(source_thread, "ready")
    repo.append_if_ready(
        source_thread,
        message_inputs(
            [HumanMessage(content="我想开一家咖啡店，目前预算有限。", id="source-message")]
        ),
    )
    result = await client.post(
        "/api/brainstorm/sessions", headers=auth, json={"topic": "咖啡店怎么获得差异化优势"}
    )
    assert result.status_code == 201, result.text
    state = result.json()
    tid = state["thread_id"]
    assert any(s["kind"] == "history" for s in state["context"]["sources"])
    excluded = [s["id"] for s in state["context"]["sources"]]
    controlled = await client.post(
        f"/api/brainstorm/sessions/{tid}/control",
        headers=auth,
        json={"action": "sources", "excluded_source_ids": excluded},
    )
    assert controlled.status_code == 200, controlled.text
    assert controlled.json()["status"] == "paused"
    assert not controlled.json()["context"]["sources"]
    assert not service.can_dispatch(state["team_id"], tid, user.id)
    payload = {"thread_id": tid, "user": str(user.id), "source": "inbox"}
    assert not await service.prepare_turn(state["team_id"], payload)
    await client.post(
        f"/api/brainstorm/sessions/{tid}/control", headers=auth, json={"action": "resume"}
    )
    assert not service.can_dispatch(state["team_id"], tid, user.id)
    await confirm_agenda(client, registry, auth, state)
    assert service.can_dispatch(state["team_id"], tid, user.id)
    service._states.clear()
    restored = await client.get(f"/api/brainstorm/sessions/{tid}", headers=auth)
    assert restored.json()["status"] == "interrupted"
    assert restored.json()["excluded_source_ids"] == excluded
    await client.post(
        f"/api/brainstorm/sessions/{tid}/control", headers=auth, json={"action": "end"}
    )
    assert (
        await client.post(
            f"/api/brainstorm/sessions/{tid}/control", headers=auth, json={"action": "resume"}
        )
    ).status_code == 400
    outsider = await grant_product(client, server, admin, "builder", "super-advisory")
    assert (
        await client.get(f"/api/brainstorm/sessions/{tid}", headers=outsider)
    ).status_code == 404


async def test_queued_native_dispatch_is_gated_and_tracked(advisory: Any, monkeypatch: Any) -> None:
    client, server, _admin, auth = advisory
    profile = (await client.post("/api/brainstorm/profile", headers=auth, json={})).json()
    registry = server.app_runtime.agent_registry
    service = service_for(registry)
    monkeypatch.setattr(service, "_launch", lambda *args: None)
    state = (
        await client.post(
            "/api/brainstorm/sessions",
            headers=auth,
            json={"topic": "新产品想法", "context_enabled": False},
        )
    ).json()
    tid = state["thread_id"]
    assert state["phase"] == "intake"
    assert not service.can_dispatch(state["team_id"], tid, state["user_id"])
    state = await confirm_agenda(client, registry, auth, state)
    from types import SimpleNamespace

    from octop.infra.agents.runtime_manager import ApplicationTeamManager

    for row in registry._repos.agent_repo.list_all():
        entry = registry._harness_manager.get_agent(row.agent_id)
        entry.metadata.update(
            user_id=row.user_id,
            application_id=row.application_id,
            teams_enabled=registry._application_teams_enabled(row),
        )
    team = ApplicationTeamManager(SimpleNamespace(get=registry._harness_manager.get_agent))
    team.set_processor(registry._team_processor)
    registry._team_processor.install_host_dispatch(team)
    monkeypatch.setattr(team._inbox, "enqueue", lambda **kwargs: "synthetic-job")
    result = await team.call_peer(
        from_agent_id=profile["team_id"],
        to_agent_id=profile["members"][0]["agent_id"],
        message="提出新方向",
        user_id=state["user_id"],
        source_thread_id=tid,
    )
    assert result.job_id == "synthetic-job"
    with pytest.raises(OctopError):
        team.submit_peer(
            from_agent_id=profile["team_id"],
            to_agent_id=profile["members"][1]["agent_id"],
            message="不能同时邀请第二位",
            user_id=state["user_id"],
            source_thread_id=tid,
        )
    await service.finish_turn(tid)
    assert (await client.get(f"/api/brainstorm/sessions/{tid}", headers=auth)).json()[
        "status"
    ] == "running"
    await client.post(
        f"/api/brainstorm/sessions/{tid}/control", headers=auth, json={"action": "pause"}
    )
    with pytest.raises(OctopError):
        await team.call_peer(
            from_agent_id=profile["team_id"],
            to_agent_id=profile["members"][0]["agent_id"],
            message="被暂停的派工",
            user_id=state["user_id"],
            source_thread_id=tid,
        )
    with pytest.raises(ValueError):
        team._resolve_peer_entry(profile["team_id"], "main", state["user_id"])


async def test_user_agenda_native_batch_deliberation_and_final_ruling(
    tmp_octop_home: Any, monkeypatch: Any
) -> None:
    """Native request/inbox/room; scripted models verify the complete meeting protocol."""
    import asyncio

    from octop.infra.agents.teams.brainstorm import build_brainstorm_tools
    from tests.support.app import octop_client
    from tests.support.auth import auth_header, bootstrap_admin, create_user, seed_openai_provider

    async with octop_client(tmp_octop_home, patch_llm=False) as (client, server):
        await bootstrap_admin(client, tmp_octop_home)
        admin = await auth_header(client)
        await seed_openai_provider(client, admin)
        await create_user(client, admin, username="reader")
        result = await client.post(
            "/api/admin/applications",
            headers=admin,
            json={
                "application_id": "super-advisory",
                "name": "超级参谋团",
                "capabilities": ["chat", "teams", "brainstorm", "zhijun_context"],
                "agent_template": {"default_model": "openai/gpt-4o"},
            },
        )
        assert result.status_code == 201
        auth = await grant_product(client, server, admin, "reader", "super-advisory")
        profile = (await client.post("/api/brainstorm/profile", headers=auth, json={})).json()
        registry = server.app_runtime.agent_registry
        service = service_for(registry)
        team = registry._harness_manager._team
        checkpoint = next(
            t
            for t in build_brainstorm_tools(registry, profile["team_id"])
            if t.name == "brainstorm_checkpoint"
        )
        host_calls = 0
        stages = {i: asyncio.Event() for i in [1, 2, 6]}
        member_release = asyncio.Event()

        async def scripted_stream(agent_id: str, payload: dict[str, Any]) -> Any:
            nonlocal host_calls
            from octop_harness.request import ChatRequest

            request = ChatRequest.coerce(payload)
            normalized = request.normalize_messages()
            runtime = request.to_runnable_config()["configurable"]
            if agent_id == profile["team_id"]:
                host_calls += 1
                state = service._states[payload["thread_id"]]
                args: dict[str, Any] = {
                    "summary": "围绕用户的时间限制寻找路线",
                    "ideas": [],
                    "directions": [],
                    "questions": [],
                }
                if host_calls == 1:
                    assert state.phase == "intake" and not state.agenda_confirmed
                    with pytest.raises(OctopError):
                        team.submit_peer(
                            from_agent_id=agent_id,
                            to_agent_id=profile["members"][0]["agent_id"],
                            message="提前派工",
                            user_id=state.user_id,
                            source_thread_id=state.thread_id,
                        )
                    args["decision"] = {
                        "kind": "intake",
                        "question": "你最想解决哪件事？",
                        "options": [{"label": "减轻我的投入"}, {"label": "扩大招生"}],
                    }
                elif host_calls == 2:
                    assert state.judgments[-1].answer == "减轻我的投入"
                    args.update(
                        phase="agenda",
                        goal="减轻我的投入",
                        constraints=["每周只有两小时"],
                        agenda="比较两条低投入路线",
                        brief="先比较投入与用户价值，由你确认后开始。",
                    )
                elif host_calls == 3:
                    assert state.agenda_confirmed and state.phase == "deliberation"
                    team.submit_peer(
                        from_agent_id=agent_id,
                        to_agent_id=profile["members"][0]["agent_id"],
                        message="按用户每周两小时的条件比较低投入路线",
                        user_id=state.user_id,
                        source_thread_id=state.thread_id,
                        metadata={"session_key": runtime["session_key"]},
                    )
                    with pytest.raises(OctopError):
                        team.submit_peer(
                            from_agent_id=agent_id,
                            to_agent_id=profile["members"][2]["agent_id"],
                            message="超出本批范围",
                            user_id=state.user_id,
                            source_thread_id=state.thread_id,
                        )
                    args["brief"] = "正在比较两条低投入路线，你可以旁听或随时发言。"
                    member_release.set()
                elif host_calls == 4:
                    incoming = "\n".join(
                        str(m.content) for m in normalized if isinstance(m, HumanMessage)
                    )
                    assert profile["members"][0]["name"] in incoming
                    assert "每周两小时" in incoming
                    team.submit_peer(
                        from_agent_id=agent_id,
                        to_agent_id=profile["members"][1]["agent_id"],
                        message="回应上一位复用课程的建议，指出具体代价",
                        user_id=state.user_id,
                        source_thread_id=state.thread_id,
                        metadata={"session_key": runtime["session_key"]},
                    )
                    args["brief"] = "先听听复用课程会不会带来新的答疑负担。"
                elif host_calls == 5:
                    incoming = "\n".join(
                        str(m.content) for m in normalized if isinstance(m, HumanMessage)
                    )
                    assert profile["members"][1]["name"] in incoming
                    team.submit_peer(
                        from_agent_id=agent_id,
                        to_agent_id=profile["members"][3]["agent_id"],
                        message="检验这两条路线在每周两小时约束下的风险",
                        user_id=state.user_id,
                        source_thread_id=state.thread_id,
                        metadata={"session_key": runtime["session_key"]},
                    )
                    args["brief"] = "两位参谋观点已汇齐，接下来检验风险。"
                else:
                    assert host_calls == 6
                    args.update(
                        brief="方案A投入较少，方案B增长可能较快但需更多投入。我推荐先验证A。",
                        decision={
                            "kind": "verdict",
                            "question": "你希望采纳哪条路线？",
                            "options": [
                                {
                                    "label": "采纳低投入方案A",
                                    "detail": "先验证两周",
                                    "effect": "adopt",
                                },
                                {"label": "继续比较", "effect": "revisit"},
                            ],
                        },
                    )
                await checkpoint.ainvoke(
                    args,
                    config={
                        "configurable": {
                            "thread_id": state.thread_id,
                            "agent_id": agent_id,
                            "user": str(state.user_id),
                            "brainstorm_revision": state.revision,
                        }
                    },
                )
                yield {"type": "token", "content": f"主持人第 {host_calls} 次呈报。"}
                if host_calls in stages:
                    stages[host_calls].set()
            else:
                await member_release.wait()
                if agent_id == profile["members"][3]["agent_id"]:
                    heard = "\n".join(str(m.content) for m in normalized)
                    assert "先复用旧课程" in heard and "复用也可能低估运营成本" in heard
                    speech = "我接着这个分歧：复用内容省了制作时间，却未必省下答疑时间。先记录两小时里实际花在哪，再决定是否继续。"
                elif agent_id == profile["members"][0]["agent_id"]:
                    speech = "每周两小时的条件下，我提一个假设：先复用旧课程，把新增内容暂停，看看学员是否仍有价值感。"
                else:
                    speech = "我补充并质疑刚才的假设：复用也可能低估运营成本。每周两小时里，课后答疑也要计入，我们能先试一次吗？"
                assert "每周两小时" in str(payload["messages"]) or "两小时" in str(
                    payload["messages"]
                )
                yield {"type": "token", "content": speech}
            yield {"type": "done"}

        monkeypatch.setattr(registry._harness_manager, "stream", scripted_stream)
        started = await client.post(
            "/api/brainstorm/sessions",
            headers=auth,
            json={"topic": "学院运营", "context_enabled": False},
        )
        assert started.status_code == 201, started.text
        tid = started.json()["thread_id"]
        await asyncio.wait_for(stages[1].wait(), 10)
        first = (await client.get(f"/api/brainstorm/sessions/{tid}", headers=auth)).json()
        assert host_calls == 1 and first["status"] == "waiting"
        response = await client.post(
            f"/api/brainstorm/sessions/{tid}/control",
            headers=auth,
            json={
                "action": "respond",
                "decision_id": first["decision"]["id"],
                "option_index": 0,
                "revision": first["revision"],
            },
        )
        assert response.status_code == 200, response.text
        await asyncio.wait_for(stages[2].wait(), 10)
        agenda = (await client.get(f"/api/brainstorm/sessions/{tid}", headers=auth)).json()
        assert agenda["phase"] == "agenda" and not agenda["agenda_confirmed"]
        assert not service.can_dispatch(profile["team_id"], tid, agenda["user_id"])
        response = await client.post(
            f"/api/brainstorm/sessions/{tid}/control",
            headers=auth,
            json={
                "action": "start",
                "decision_id": agenda["decision"]["id"],
                "revision": agenda["revision"],
            },
        )
        assert response.status_code == 200, response.text
        await asyncio.wait_for(stages[6].wait(), 15)
        await asyncio.sleep(0.15)
        final = (await client.get(f"/api/brainstorm/sessions/{tid}", headers=auth)).json()
        assert host_calls == 6 and final["phase"] == "decision" and final["status"] == "waiting"
        assert final["goal"] == "减轻我的投入" and final["turns"] == 6
        response = await client.post(
            f"/api/brainstorm/sessions/{tid}/control",
            headers=auth,
            json={
                "action": "respond",
                "decision_id": final["decision"]["id"],
                "option_index": 0,
                "revision": final["revision"],
            },
        )
        assert response.status_code == 200, response.text
        assert response.json()["status"] == "ended" and response.json()["ruling"].startswith(
            "采纳低投入方案A"
        )
        repeat = await client.post(
            f"/api/brainstorm/sessions/{tid}/control",
            headers=auth,
            json={"action": "respond", "decision_id": final["decision"]["id"], "option_index": 0},
        )
        assert repeat.status_code == 400
        messages, _ = server.services.thread_message_repo.page(tid, limit=100, offset=0)
        assert any("每周两小时" in m.message_json for m in messages)
        assert any("减轻我的投入" in m.message_json and m.role == "human" for m in messages)


async def test_host_failure_is_visible_and_does_not_count_progress(
    advisory: Any, monkeypatch: Any
) -> None:
    import asyncio

    client, server, _admin, auth = advisory
    await client.post("/api/brainstorm/profile", headers=auth, json={})
    registry = server.app_runtime.agent_registry
    frames = []
    visible = asyncio.Event()

    async def failed_stream(*args: Any, **kwargs: Any) -> Any:
        raise RuntimeError("scripted model failure")
        yield  # pragma: no cover

    async def push_to_thread(thread_id: str, frame: Any) -> None:
        frames.append(frame)
        if frame.get("type") == "error":
            visible.set()

    monkeypatch.setattr(registry._harness_manager, "stream", failed_stream)
    monkeypatch.setattr(registry._team_processor._gateway.ws_hub, "push_to_thread", push_to_thread)
    result = await client.post(
        "/api/brainstorm/sessions",
        headers=auth,
        json={"topic": "失败后的恢复", "context_enabled": False},
    )
    assert result.status_code == 201, result.text
    await asyncio.wait_for(visible.wait(), timeout=5)
    state = (
        await client.get(f"/api/brainstorm/sessions/{result.json()['thread_id']}", headers=auth)
    ).json()
    assert state["status"] == "interrupted" and state["turns"] == 0
    assert any(
        frame.get("type") == "error" and frame["message"] in state["context"]["notices"]
        for frame in frames
    )


async def test_feedback_stale_cards_pause_and_scope_reconfirmation(
    advisory: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    from types import SimpleNamespace

    from octop.infra.agents.teams.brainstorm import build_brainstorm_tools

    client, server, _admin, auth = advisory
    profile = (await client.post("/api/brainstorm/profile", headers=auth, json={})).json()
    registry = server.app_runtime.agent_registry
    service = service_for(registry)
    monkeypatch.setattr(service, "_launch", lambda *_: None)
    state = (
        await client.post(
            "/api/brainstorm/sessions",
            headers=auth,
            json={"topic": "我希望减少运营投入", "context_enabled": False},
        )
    ).json()
    tid = state["thread_id"]
    checkpoint = next(
        t
        for t in build_brainstorm_tools(registry, profile["team_id"])
        if t.name == "brainstorm_checkpoint"
    )

    async def save(**kwargs: Any) -> None:
        await checkpoint.ainvoke(
            {"summary": "秘书呈报", "ideas": [], "directions": [], "questions": [], **kwargs},
            config={
                "configurable": {
                    "thread_id": tid,
                    "agent_id": profile["team_id"],
                    "user": str(state["user_id"]),
                    "brainstorm_revision": service._states[tid].revision,
                }
            },
        )

    async def control(**kwargs: Any) -> Any:
        response = await client.post(
            f"/api/brainstorm/sessions/{tid}/control",
            headers=auth,
            json={"revision": service._states[tid].revision, **kwargs},
        )
        assert response.status_code == 200, response.text
        return response.json()

    await save(
        decision={
            "kind": "intake",
            "question": "哪项限制最重要？",
            "options": [{"label": "我的时间"}, {"label": "预算"}],
        }
    )
    first_card = service._states[tid].decision
    assert first_card is not None
    resumed = await control(action="resume")
    assert resumed["status"] == "waiting" and resumed["decision"]["id"] == first_card.id
    answered = await control(
        action="respond",
        decision_id=first_card.id,
        text="我每周只能投入两小时，而且不想再增加人员。",
    )
    assert answered["judgments"][-1]["answer"].startswith("我每周只能投入两小时")
    assert not service.can_dispatch(profile["team_id"], tid, state["user_id"])
    snapshot = service._states[tid].model_dump()
    with pytest.raises(OctopError):
        await save(
            brief="不应写入的内容",
            decision={
                "kind": "verdict",
                "question": "采纳？",
                "options": [
                    {"label": "采纳", "effect": "adopt"},
                    {"label": "继续", "effect": "revisit"},
                ],
            },
        )
    assert service._states[tid].model_dump() == snapshot
    stale = await client.post(
        f"/api/brainstorm/sessions/{tid}/control",
        headers=auth,
        json={"action": "respond", "decision_id": first_card.id, "option_index": 0},
    )
    assert stale.status_code == 400
    stale = await client.post(
        f"/api/brainstorm/sessions/{tid}/control",
        headers=auth,
        json={"action": "respond", "text": "旧页面的反馈", "revision": 0},
    )
    assert stale.status_code == 400
    started = await confirm_agenda(client, registry, auth, answered)
    service.track_job(tid, "old-job", begin=True)
    old_revision = started["revision"]
    await control(action="pause")
    old = SimpleNamespace(
        id="old-job",
        source_thread_id=tid,
        source_agent_id=profile["team_id"],
        target_agent_id=profile["members"][0]["agent_id"],
    )
    assert await service.collect_member_result(old, "旧方向的建议", None) is None
    assert service._states[tid].status == "paused"
    with pytest.raises(OctopError):
        await checkpoint.ainvoke(
            {"summary": "过期总结", "ideas": [], "directions": [], "questions": []},
            config={
                "configurable": {
                    "thread_id": tid,
                    "agent_id": profile["team_id"],
                    "user": str(state["user_id"]),
                    "brainstorm_revision": old_revision,
                }
            },
        )
    await control(action="resume")
    for job in ("one", "two"):
        service.track_job(tid, job, begin=True)
        service.track_job(tid, job, begin=False)
    one = SimpleNamespace(**{**vars(old), "id": "one"})
    two = SimpleNamespace(
        **{**vars(old), "id": "two", "target_agent_id": profile["members"][1]["agent_id"]}
    )
    assert await service.collect_member_result(one, "路线A", None) is None
    collected = await service.collect_member_result(two, "路线B", None)
    assert collected is not None and "路线A" in collected[0] and "路线B" in collected[0]
    assert await service.collect_member_result(two, "重复回复", None) is None
    await save(
        decision={
            "kind": "judgment",
            "question": "是否可以增加投入？",
            "options": [{"label": "保持两小时"}, {"label": "可以增加"}],
        }
    )
    card = service._states[tid].decision
    assert card is not None
    await control(action="resume")
    assert service._states[tid].status == "waiting"
    await control(action="respond", decision_id=card.id, text="仍然保持两小时，请继续比较。")
    assert service.can_dispatch(profile["team_id"], tid, state["user_id"])
    await save(
        decision={
            "kind": "verdict",
            "question": "请选择路线",
            "options": [
                {"label": "采纳A", "effect": "adopt"},
                {"label": "重新比较", "effect": "revisit"},
            ],
        }
    )
    card = service._states[tid].decision
    assert card is not None
    revisited = await control(action="respond", decision_id=card.id, option_index=1)
    assert revisited["phase"] == "deliberation" and revisited["ruling"] == ""
    assert revisited["judgments"][-1]["answer"] == "重新比较"
    revised = await control(action="revise")
    assert revised["phase"] == "intake" and not revised["agenda_confirmed"]
    assert not service.can_dispatch(profile["team_id"], tid, state["user_id"])
    await control(action="agenda")
    await save(agenda="先验证留存，不再讨论扩招")
    assert service._states[tid].decision is not None
    assert service._states[tid].decision.kind == "agenda"
    assert not service.can_dispatch(profile["team_id"], tid, state["user_id"])

    refreshed = (await client.get(f"/api/brainstorm/sessions/{tid}", headers=auth)).json()
    # Written ruling is explicit, and is only permitted on a verdict card.
    invalid = await client.post(
        f"/api/brainstorm/sessions/{tid}/control",
        headers=auth,
        json={
            "action": "adopt",
            "text": "采纳我的修改",
            "decision_id": refreshed["decision"]["id"],
        },
    )
    assert invalid.status_code == 400
    await control(action="start", decision_id=refreshed["decision"]["id"])
    await save(
        decision={
            "kind": "verdict",
            "question": "请选择路线",
            "options": [
                {"label": "采纳A", "effect": "adopt"},
                {"label": "重新比较", "effect": "revisit"},
            ],
        }
    )
    card = service._states[tid].decision
    assert card is not None
    written = await control(
        action="adopt", decision_id=card.id, text="采纳A，但先做一周验证，每周投入不超过两小时。"
    )
    assert written["phase"] == "closed" and written["status"] == "ended"
    assert written["ruling"] == "采纳A，但先做一周验证，每周投入不超过两小时。"
    assert written["report"]["outcome"] == "adopted"
    assert written["ruling"] in written["report"]["markdown"]


@pytest.mark.parametrize("adopt", [False, True])
async def test_report_freezes_owner_ruling_and_exports_real_docx(
    advisory: Any, monkeypatch: pytest.MonkeyPatch, adopt: bool
) -> None:
    from io import BytesIO

    from docx import Document
    from langchain_core.messages import AIMessage

    from octop.infra.agents.teams.brainstorm import build_brainstorm_tools
    from octop.infra.history.projection import message_inputs

    client, server, admin, auth = advisory
    registry = server.app_runtime.agent_registry
    service = service_for(registry)
    monkeypatch.setattr(service, "_launch", lambda *_: None)
    await client.post("/api/brainstorm/profile", headers=auth, json={})
    state = (
        await client.post(
            "/api/brainstorm/sessions",
            headers=auth,
            json={"topic": "学院低投入运营", "context_enabled": False},
        )
    ).json()
    tid = state["thread_id"]
    for suffix in ["report", "report/document"]:
        assert (
            await client.get(f"/api/brainstorm/sessions/{tid}/{suffix}", headers=auth)
        ).status_code == 400
    await confirm_agenda(client, registry, auth, state)
    checkpoint = next(
        t
        for t in build_brainstorm_tools(registry, state["team_id"])
        if t.name == "brainstorm_checkpoint"
    )
    await checkpoint.ainvoke(
        {
            "summary": "先小规模验证",
            "ideas": ["两条低投入路线"],
            "directions": [],
            "questions": [],
            "report_draft": {
                "overview": "在每周两小时内验证学员留存。",
                "key_points": ["复用课程素材"],
                "disagreements": ["战略参谋主张筛选，实践参谋强调承接。"],
                "actions": [{"task": "回访五名学员"}],
                "open_questions": ["是否收费仍待确认"],
            },
            "decision": {
                "kind": "verdict",
                "question": "如何推进？",
                "options": [
                    {"label": "采纳小规模验证", "effect": "adopt"},
                    {"label": "继续研究", "effect": "revisit"},
                ],
            },
        },
        config={
            "configurable": {
                "thread_id": tid,
                "agent_id": state["team_id"],
                "user": str(state["user_id"]),
            }
        },
    )
    repo = server.services.thread_message_repo
    member = state["participants"][0]
    repo.append_if_ready(
        tid,
        message_inputs(
            [
                AIMessage(
                    content="先验证留存，再决定扩招。",
                    id="before-close",
                    additional_kwargs={"speaker_agent_id": member["agent_id"]},
                )
            ]
        ),
    )
    card = service._states[tid].decision
    assert card is not None
    body = (
        {"action": "adopt", "decision_id": card.id, "text": "先验证一周，每周最多两小时。"}
        if adopt
        else {"action": "end"}
    )
    closed = await client.post(f"/api/brainstorm/sessions/{tid}/control", headers=auth, json=body)
    assert closed.status_code == 200, closed.text
    report = closed.json()["report"]
    assert report["outcome"] == ("adopted" if adopt else "unresolved")
    assert "先验证一周" in report["markdown"] if adopt else "尚未作出最终裁决" in report["markdown"]
    for text in [
        "复用课程素材",
        "战略参谋主张筛选",
        "回访五名学员",
        "是否收费仍待确认",
        "先验证留存",
        "待确认",
    ]:
        assert text in report["markdown"]
    repo.append_if_ready(
        tid,
        message_inputs(
            [
                AIMessage(
                    content="迟到的结果不可改写报告",
                    id="after-close",
                    additional_kwargs={"speaker_agent_id": member["agent_id"]},
                )
            ]
        ),
    )
    service._states.clear()
    assert (
        await client.get(f"/api/brainstorm/sessions/{tid}/report", headers=auth)
    ).json() == report
    exported = await client.get(f"/api/brainstorm/sessions/{tid}/report/document", headers=auth)
    assert exported.status_code == 200, exported.text[:200] if exported.status_code != 200 else ""
    assert (
        exported.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    document = Document(BytesIO(exported.content))
    assert document.paragraphs[0].style.name == "Title"
    text = "\n".join(p.text for p in document.paragraphs)
    assert "回访五名学员" in text and "迟到的结果" not in text
    other = await grant_product(client, server, admin, "builder", "super-advisory")
    for suffix in ["report", "report/document"]:
        assert (
            await client.get(f"/api/brainstorm/sessions/{tid}/{suffix}", headers=other)
        ).status_code == 404
    await client.put(
        "/api/admin/applications/super-advisory/access",
        headers=admin,
        json={"user_id": state["user_id"], "enabled": False},
    )
    assert (
        await client.get(f"/api/brainstorm/sessions/{tid}/report/document", headers=auth)
    ).status_code == 403


async def test_legacy_meeting_requires_owner_agenda_confirmation(
    advisory: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    import json

    client, server, _admin, auth = advisory
    profile = (await client.post("/api/brainstorm/profile", headers=auth, json={})).json()
    registry = server.app_runtime.agent_registry
    service = service_for(registry)
    monkeypatch.setattr(service, "_launch", lambda *_: None)
    state = (
        await client.post(
            "/api/brainstorm/sessions",
            headers=auth,
            json={"topic": "历史会议", "context_enabled": False},
        )
    ).json()
    tid = state["thread_id"]
    for key in (
        "phase",
        "agenda_confirmed",
        "goal",
        "constraints",
        "agenda",
        "brief",
        "decision",
        "judgments",
        "intake_answers",
        "ruling",
    ):
        state.pop(key)
    await service._workspace(profile["team_id"]).aupload_bytes(
        f"brainstorm/{tid}.json", json.dumps(state).encode()
    )
    service._states.clear()
    restored = (await client.get(f"/api/brainstorm/sessions/{tid}", headers=auth)).json()
    assert restored["phase"] == "intake" and restored["status"] == "interrupted"
    assert restored["goal"] == "历史会议" and not restored["agenda_confirmed"]
    assert not service.can_dispatch(profile["team_id"], tid, restored["user_id"])


async def test_flexible_experts_and_meeting_roster_isolation(
    advisory: Any, monkeypatch: Any
) -> None:
    client, server, admin, auth = advisory
    registry = server.app_runtime.agent_registry
    service = service_for(registry)
    monkeypatch.setattr(service, "_launch", lambda *args: None)
    ids = ["brainstorm-product", "brainstorm-user", "brainstorm-organization"]
    response = await client.post("/api/brainstorm/profile", headers=auth, json={"expert_ids": ids})
    assert response.status_code == 200, response.text
    profile = response.json()
    assert {m["expert_id"] for m in profile["members"]} == set(ids)
    assert sum(e["recommended"] for e in profile["experts"]) == 8
    # Host deployment bindings and positional legacy model writes stay out of scope.
    for body in ({"host_model": "openai/gpt-4o"}, {"member_models": ["openai/gpt-4o"] * 4}):
        assert (
            await client.post("/api/brainstorm/profile", headers=auth, json=body)
        ).status_code == 422
    assert (await client.post("/api/admin/providers", headers=auth, json={})).status_code == 404
    assert (await client.get("/api/admin/providers", headers=auth)).status_code == 404
    assert (
        await client.post(
            "/api/brainstorm/profile", headers=auth, json={"expert_ids": ["missing", ids[0]]}
        )
    ).status_code == 400
    state = (
        await client.post(
            "/api/brainstorm/sessions",
            headers=auth,
            json={"topic": "研究组织协作", "context_enabled": False, "expert_ids": ids[:2]},
        )
    ).json()
    assert {m["expert_id"] for m in state["participants"]} == set(ids[:2])
    state = await confirm_agenda(client, registry, auth, state)
    chosen = state["participants"][0]["agent_id"]
    excluded = next(m["agent_id"] for m in profile["members"] if m["expert_id"] == ids[2])
    processor = registry._team_processor
    processor._brainstorm_dispatch_guard(
        profile["team_id"], state["thread_id"], state["user_id"], chosen
    )
    with pytest.raises(OctopError):
        processor._brainstorm_dispatch_guard(
            profile["team_id"], state["thread_id"], state["user_id"], excluded
        )
    await client.post(
        f"/api/brainstorm/sessions/{state['thread_id']}/control",
        headers=auth,
        json={"action": "pause"},
    )
    next_ids = ["brainstorm-business", "brainstorm-creative"]
    updated = await client.post(
        "/api/brainstorm/profile", headers=auth, json={"expert_ids": next_ids}
    )
    assert updated.status_code == 200, updated.text
    assert {m["expert_id"] for m in updated.json()["members"]} == set(next_ids)
    # The paused meeting still retains and can dispatch its original members.
    resumed = await client.post(
        f"/api/brainstorm/sessions/{state['thread_id']}/control",
        headers=auth,
        json={"action": "resume"},
    )
    assert resumed.json()["participants"] == state["participants"]
    assert chosen in registry.teams.member_ids(profile["team_id"])
    assert service.can_dispatch(profile["team_id"], state["thread_id"], state["user_id"], chosen)
    assert not service.can_dispatch(
        profile["team_id"], state["thread_id"], state["user_id"], excluded
    )


async def test_expert_models_use_wanx_catalog_and_persist_native_assignments(
    advisory_native: Any, monkeypatch: Any
) -> None:
    client, server, admin, auth = advisory_native
    registry = server.app_runtime.agent_registry
    service = service_for(registry)
    monkeypatch.setattr(service, "_launch", lambda *args: None)
    provider = await client.post(
        "/api/admin/providers",
        headers=admin,
        json={
            "name": "advisory-test",
            "kind": "openai",
            "base_url": "https://example.invalid/v1",
            "api_key": "synthetic-key",
            "models": [
                {"id": "chat", "enabled": True},
                {"id": "disabled", "enabled": False},
                {"id": "embedding", "enabled": True, "embedding": True},
            ],
        },
    )
    assert provider.status_code == 201, provider.text
    ids = ["brainstorm-product", "brainstorm-challenge"]
    assignments = dict(zip(ids, ["advisory-test/chat", "openai/gpt-4o"], strict=True))
    catalog = (await client.get("/api/brainstorm/profile", headers=auth)).json()
    assert set(catalog["models"]) == set(assignments.values())
    # Reject the entire request before creating any native agents.
    for invalid in (
        {ids[0]: "missing/chat"},
        {ids[0]: "advisory-test/disabled"},
        {ids[0]: "advisory-test/embedding"},
        {"brainstorm-creative": "openai/gpt-4o"},
    ):
        rejected = await client.post(
            "/api/brainstorm/profile",
            headers=auth,
            json={"expert_ids": ids, "expert_models": invalid},
        )
        assert rejected.status_code == 400, rejected.text
        assert (await client.get("/api/brainstorm/profile", headers=auth)).json()["team_id"] is None
    response = await client.post(
        "/api/brainstorm/profile",
        headers=auth,
        json={"expert_ids": ids, "expert_models": assignments},
    )
    assert response.status_code == 200, response.text
    profile = response.json()
    assert {m["expert_id"]: m["model"] for m in profile["members"]} == assignments
    assert profile["host_model"] == "openai/gpt-4o"
    members = {m["expert_id"]: m["agent_id"] for m in profile["members"]}
    updated = await client.post(
        "/api/brainstorm/profile",
        headers=auth,
        json={"expert_models": {ids[1]: "advisory-test/chat"}},
    )
    assert updated.status_code == 200, updated.text
    readback = (await client.get("/api/brainstorm/profile", headers=auth)).json()
    assert {m["expert_id"]: m["agent_id"] for m in readback["members"]} == members
    assert all(m["model"] == "advisory-test/chat" for m in readback["members"])
    for aid in members.values():
        await registry.reload(aid)
        assert (
            registry._providers.resolve_explicit_default_model(
                registry.get_row(aid), registry.get_config(aid)
            )
            == "advisory-test/chat"
        )
    # A model disabled after the UI loaded cannot be saved from that stale catalog.
    disabled = await client.patch(
        f"/api/admin/providers/{provider.json()['id']}", headers=admin, json={"enabled": False}
    )
    assert disabled.status_code == 200, disabled.text
    unavailable = (await client.get("/api/brainstorm/profile", headers=auth)).json()
    assert unavailable["models"] == ["openai/gpt-4o"]
    assert unavailable["members"] == readback["members"]
    rejected = await client.post(
        "/api/brainstorm/profile", headers=auth, json={"expert_models": assignments}
    )
    assert rejected.status_code == 400
    assert (await client.get("/api/brainstorm/profile", headers=auth)).json()[
        "members"
    ] == readback["members"]


async def test_model_assignment_cannot_change_during_deliberation(
    advisory: Any, monkeypatch: Any
) -> None:
    client, server, _, auth = advisory
    registry = server.app_runtime.agent_registry
    monkeypatch.setattr(service_for(registry), "_launch", lambda *args: None)
    profile = (await client.post("/api/brainstorm/profile", headers=auth, json={})).json()
    state = (
        await client.post(
            "/api/brainstorm/sessions",
            headers=auth,
            json={"topic": "验证模型配置锁", "context_enabled": False},
        )
    ).json()
    state = await confirm_agenda(client, registry, auth, state)
    assert state["status"] == "running"
    busy = await client.post(
        "/api/brainstorm/profile",
        headers=auth,
        json={"expert_models": {profile["members"][0]["expert_id"]: "openai/gpt-4o"}},
    )
    assert busy.status_code == 409
    assert busy.json()["error"]["code"] == "TEAM_MEMBER_BUSY"
    assert (await client.get("/api/brainstorm/profile", headers=auth)).json()["members"] == profile[
        "members"
    ]


async def test_skillhub_import_keeps_native_workflow_and_skills(
    advisory_native: Any, monkeypatch: Any
) -> None:
    import io
    import json
    import zipfile

    from octop.infra.agents.experts import skillhub_market as market
    from octop.infra.skills import skillhub_market as skills

    client, server, _, auth = advisory_native
    registry = server.app_runtime.agent_registry
    item = market.SkillHubSkillset(
        slug="growth-lab",
        display_name="增长实验专家",
        summary="检验增长假设",
        skill_slugs=("growth-test",),
    )
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr("manifest.json", json.dumps({"skillSlugs": ["growth-test"]}))
        z.writestr("skillsets/growth-lab.md", "# 增长工作流\n先检验假设，再看证据。")
    monkeypatch.setattr(market, "fetch_skillset", lambda slug: item)
    monkeypatch.setattr(market, "_download_skillset_package", lambda slug: archive.getvalue())
    monkeypatch.setattr(
        skills,
        "download_skillhub_package_files",
        lambda slug: [
            ("SKILL.md", b"---\nname: growth-test\ndescription: test\n---\nTest a hypothesis.")
        ],
    )
    monkeypatch.setattr(
        "octop.api.routers.experts.browse_skillsets", lambda *a, **kw: ([item], ["marketing"])
    )
    response = await client.get("/api/experts/hub?q=growth", headers=auth)
    assert response.status_code == 200, response.text
    assert response.json()["items"][0]["id"] == item.expert_id
    imported = await client.post("/api/brainstorm/experts/growth-lab/import", headers=auth)
    assert imported.status_code == 200, imported.text
    assert imported.json()["source"] == "skillhub"
    selected = await client.post(
        "/api/brainstorm/profile",
        headers=auth,
        json={"expert_ids": [item.expert_id, "brainstorm-challenge"]},
    )
    assert selected.status_code == 200, selected.text
    member = next(m for m in selected.json()["members"] if m["expert_id"] == item.expert_id)
    row = registry.get_row(member["agent_id"])
    assert row.application_id == "super-advisory" and not row.is_shared
    workspace = registry.workspace_for_agent(row.agent_id)
    assert "增长实验专家" in await workspace.aread_text("SOUL.md")
    assert "先检验假设" in await workspace.aread_text("skills/growth-lab/SKILL.md")
    assert "Test a hypothesis." in await workspace.aread_text("skills/growth-test/SKILL.md")
